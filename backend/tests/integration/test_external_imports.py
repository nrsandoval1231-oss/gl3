import json
from datetime import UTC, datetime, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

import pytest
from db_reset import reset_test_database
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy import event as sqlalchemy_event
from sqlalchemy.engine import make_url
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker

from gridlens.ingestion.workbook_import import (
    HYPERLINK_METADATA_PARSER_REVISION,
    ImportPreview,
    ImportRow,
    read_workbook,
)
from gridlens.intelligence.company_contact_discovery import (
    JOB_KIND as CONTACT_JOB_KIND,
)
from gridlens.intelligence.company_contact_discovery import (
    _research_source_digest,
    plan_company_contact_jobs,
)
from gridlens.intelligence.external_imports import (
    MAX_COMPANY_HYPERLINK_TARGET_BYTES,
    _is_hyperlink_sidecar,
    _metadata_sidecar_keys,
    company_hyperlink_hints,
    import_history,
    reconcile_preview,
    rollback_import,
    stage_import,
)
from gridlens.persistence.models import (
    BuyingChainEvidence,
    BuyingChainRun,
    CompanyIdentity,
    CompanyProjectMembership,
    Evidence,
    ExternalImportActivation,
    ExternalImportBatch,
    ExternalImportEvent,
    ExternalImportRow,
    IdentityIdentifier,
    Interpretation,
    Job,
    Observation,
    PowerNeed,
    Project,
    Source,
)
from gridlens.settings import settings


@pytest.fixture
def import_db():
    target = make_url(settings.test_database_url)
    application = make_url(settings.database_url)
    assert target.database and (target.database == "test" or target.database.endswith("_test")), (
        "Refusing destructive setup without an explicit *_test database"
    )
    assert (target.host, target.port, target.database) != (
        application.host,
        application.port,
        application.database,
    )
    engine = create_engine(settings.test_database_url)
    reset_test_database(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    try:
        yield factory
    finally:
        engine.dispose()


def _preview(rows: list[ImportRow], digest: str = "a" * 64) -> ImportPreview:
    return ImportPreview(
        filename="sourced.xlsx",
        content_hash=digest,
        rows=rows,
        report={"processed_rows": len(rows)},
    )


def _seed_matchable(session):
    source = Source(
        id="dallas-monthly-permits",
        publisher="Synthetic",
        endpoint="https://example.test",
        family="construction-permits",
        geography="TX",
        industries=[],
        retrieval_method="dallas-monthly-permits",
    )
    session.add(source)
    session.flush()
    project = Project(
        identity_key="test-site-1",
        name="North Campus",
        location="Austin, TX",
        source_id=source.id,
        external_id="source-project-1",
    )
    observation = Observation(
        source_id=source.id,
        external_id="evidence-1",
        title="North Campus construction permit",
        abstract=(
            "North Campus project site location: Austin, TX. Synthetic identity corroboration."
        ),
        url="https://example.test/1",
        content_hash="b" * 64,
    )
    session.add_all([project, observation])
    session.flush()
    evidence = Evidence(
        observation_id=observation.id,
        content_hash=observation.content_hash,
        title=observation.title,
        abstract=observation.abstract,
        url=observation.url,
        extraction_location="test excerpt",
        provenance={},
    )
    session.add(evidence)
    session.flush()
    retained = BuyingChainEvidence(
        project_id=project.id,
        retained_evidence_id=None,
        url=observation.url,
        title=observation.title,
        excerpt=observation.abstract,
        content_hash=observation.content_hash,
        extraction_location="test excerpt",
        published_at=None,
        observed_at=datetime.now(UTC),
        last_seen_at=datetime.now(UTC),
        project_match_reason="synthetic test seed",
    )
    identity = CompanyIdentity()
    session.add_all(
        [
            retained,
            identity,
            Interpretation(
                observation_id=observation.id,
                content_hash=observation.content_hash,
                assessment={
                    "decision": "accepted",
                    "evidence_hash": observation.content_hash,
                    "project_identity_key": project.identity_key,
                    "subject": project.name,
                    "location": project.location,
                    "power_activity": "construction",
                    "admission_status": "admitted",
                },
                status="accepted",
            ),
        ]
    )
    session.flush()
    session.add(
        BuyingChainRun(
            project_id=project.id,
            input_hash="d" * 64,
            evidence_ids=[str(retained.id)],
            claim_ids=[],
            policy_version="test",
            status="complete",
            documents_checked=1,
            candidates_added=0,
            unresolved=[],
        )
    )
    session.flush()
    session.add_all(
        [
            IdentityIdentifier(
                entity_type="company",
                company_identity_id=identity.id,
                identifier_type="domain",
                normalized_value="acme.example",
                source_evidence_id=retained.id,
                source_version_hash=observation.content_hash,
                reason="synthetic exact domain fixture",
            ),
            CompanyProjectMembership(
                company_identity_id=identity.id,
                project_id=project.id,
                role="owner",
                organization_name="Acme Energy",
                source_id=source.id,
                source_evidence_id=evidence.id,
                source_version_hash=observation.content_hash,
                decision_basis="test fixture",
                operator_number="",
            ),
        ]
    )
    session.flush()
    return project, identity


def test_stage_is_replay_safe_explainable_and_reversible_without_canonical_writes(import_db):
    with import_db.begin() as session:
        project, identity = _seed_matchable(session)
        preview = _preview(
            [
                ImportRow(
                    sheet="Accounts",
                    row_number=2,
                    raw={"Company": "Acme Energy", "Site": "North Campus"},
                    normalized={
                        "company_name": "Acme Energy",
                        "domain": "https://www.acme.example",
                        "project_name": "North Campus",
                        "city": "Austin",
                        "state": "TX",
                        "contact_name": "Unverified Contact",
                        "email": "person@acme.example",
                        "demand_mw": "12",
                        "supply_mw": None,
                    },
                    errors=[],
                ),
                ImportRow(
                    sheet="Accounts",
                    row_number=3,
                    raw={"Company": "Acme Energy"},
                    normalized={"company_name": "Acme Energy"},
                    errors=[],
                ),
                ImportRow(
                    sheet="Accounts",
                    row_number=4,
                    raw={"Company": "Acme Energy", "Site": "North Campus"},
                    normalized={
                        "company_name": "Acme Energy",
                        "project_name": "North Campus",
                        "city": "Austin",
                        "state": "TX",
                    },
                    errors=[],
                ),
            ]
        )
        first = stage_import(session, preview, actor="operator@example.test")
        replay = stage_import(session, preview, actor="operator@example.test")
        changed_mapping = _preview(
            [
                ImportRow(
                    sheet="Accounts",
                    row_number=2,
                    raw={"Company": "Acme Energy"},
                    normalized={"company_name": "Different mapped value"},
                    errors=[],
                )
            ]
        )
        with pytest.raises(ValueError, match="different normalized preview"):
            stage_import(session, changed_mapping, actor="operator@example.test")
        assert first["batch_id"] == replay["batch_id"]
        assert first["row_count"] == replay["row_count"] == 3
        assert first["matched_count"] == 2
        assert first["unmatched_count"] == 1
        assert first["rows"][0]["company_match"]["candidate_id"] == str(identity.id)
        assert first["rows"][0]["project_match"]["candidate_id"] == str(project.id)
        assert "exact registered domain" in first["rows"][0]["company_match"]["reasons"]
        assert first["rows"][1]["company_match"]["status"] == "unmatched"
        assert first["rows"][2]["company_match"]["candidate_id"] == str(identity.id)
        assert first["rows"][2]["project_match"]["candidate_id"] == str(project.id)
        assert (
            "independent exact project name and site location"
            in first["rows"][2]["company_match"]["reasons"]
        )
        assert session.scalar(select(func.count()).select_from(ExternalImportBatch)) == 1
        assert session.scalar(select(func.count()).select_from(ExternalImportRow)) == 3
        assert session.scalar(select(func.count()).select_from(Project)) == 1
        assert session.scalar(select(func.count()).select_from(PowerNeed)) == 0
        assert session.scalar(select(func.count()).select_from(Evidence)) == 1
        retained = session.scalar(
            select(ExternalImportRow).where(ExternalImportRow.row_number == 2)
        )
        assert retained.normalized["demand_mw"] == "12"
        assert retained.normalized["supply_mw"] is None
        assert retained.normalized["contact_name"] == "Unverified Contact"

        rolled = rollback_import(session, first["batch_id"], actor="operator@example.test")
        again = rollback_import(session, first["batch_id"], actor="operator@example.test")
        assert rolled["status"] == again["status"] == "rolled_back"
        assert rolled["active_count"] == again["active_count"] == 0
        assert session.scalar(select(func.count()).select_from(ExternalImportRow)) == 3
        assert session.scalar(select(func.count()).select_from(ExternalImportEvent)) == 2
        assert all(
            not activation.active
            for activation in session.scalars(select(ExternalImportActivation))
        )
        assert import_history(session)[0]["batch_id"] == first["batch_id"]


def test_same_file_parser_metadata_reparse_is_append_only_idempotent_and_unverified(
    import_db, monkeypatch
):
    with TemporaryDirectory() as temporary_directory:
        workbook_path = Path(temporary_directory) / "linked-projects.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Projects"
        repeated_reserved_header = "__workbook_hyperlinks__"
        headers = [repeated_reserved_header] * 10 + [
            "Project Name",
            "State",
            "Citation",
            "Citation Two",
            "Citation Three",
        ]
        sheet.append(headers)
        sheet.append(["raw"] * 10 + ["Unmatched Project", "TX", "Link", "Email", "Malformed"])
        sheet["M2"].hyperlink = "https://example.gov/permits/123"
        sheet["N2"].hyperlink = "mailto:permits@example.gov"
        sheet["O2"].hyperlink = "https://["
        workbook.save(workbook_path)

        current_preview = read_workbook(workbook_path)
        current_preview.rows = [
            row.model_copy(update={"raw": {**row.raw, "Irrelevant payload": "x" * 1_000_000}})
            for row in current_preview.rows
        ]
        current_preview.report = {
            **current_preview.report,
            "irrelevant_report_blob": "y" * 1_000_000,
        }
        current_sidecar_keys = _metadata_sidecar_keys(current_preview.report)
        legacy_report = dict(current_preview.report)
        legacy_report.pop("hyperlink_metadata_parser_revision")
        legacy_report.pop("hyperlink_metadata_sidecar_keys")
        legacy_rows = [
            row.model_copy(
                update={
                    "raw": {
                        key: value
                        for key, value in row.raw.items()
                        if not _is_hyperlink_sidecar(
                            key,
                            value,
                            sidecar_keys=current_sidecar_keys.get(row.sheet, set()),
                        )
                    }
                }
            )
            for row in current_preview.rows
        ]
        legacy_preview = ImportPreview(
            filename=current_preview.filename,
            content_hash=current_preview.content_hash,
            rows=legacy_rows,
            report=legacy_report,
        )

    with import_db.begin() as session:
        original = stage_import(session, legacy_preview, actor="initial-importer")
        original_row = session.scalar(select(ExternalImportRow))
        original_raw = dict(original_row.raw)
        original_normalized = dict(original_row.normalized)
        original_batch = session.get(ExternalImportBatch, original_row.batch_id)
        original_report = dict(original_batch.report)
        assert original["active_count"] == 0
        assert original_row.raw["Citation"] == "Link"
        assert not any(
            _is_hyperlink_sidecar(
                key,
                value,
                sidecar_keys=current_sidecar_keys.get(original_row.sheet, set()),
            )
            for key, value in original_row.raw.items()
        )

        reparsed = stage_import(session, current_preview, actor="metadata-reviewer")
        assert reparsed["batch_id"] == original["batch_id"]
        assert reparsed["active_count"] == 0
        assert reparsed["recovered_unverified_hyperlink_hint_count"] == 3
        receipt = reparsed["metadata_reparse_receipts"][0]
        assert receipt["actor"] == "metadata-reviewer"
        assert receipt["source_content_hash"] == current_preview.content_hash
        assert receipt["parser_revision"] == HYPERLINK_METADATA_PARSER_REVISION
        assert receipt["hyperlink_hints"] == [
            {
                "sheet": "Projects",
                "row_number": 2,
                "column": "Citation",
                "url": "https://example.gov/permits/123",
                "target": "https://example.gov/permits/123",
                "verification_status": "UNVERIFIED_SOURCE_HINT",
            },
            {
                "sheet": "Projects",
                "row_number": 2,
                "column": "Citation Two",
                "url": None,
                "target": "mailto:permits@example.gov",
                "verification_status": "UNVERIFIED_UNSUPPORTED_HYPERLINK_TARGET",
            },
            {
                "sheet": "Projects",
                "row_number": 2,
                "column": "Citation Three",
                "url": None,
                "target": "https://[",
                "verification_status": "UNVERIFIED_UNSUPPORTED_HYPERLINK_TARGET",
            },
        ]
        assert reparsed["rows"][0]["unverified_hyperlink_hints"][0]["url"] == (
            "https://example.gov/permits/123"
        )

        # The shared reader selects only the current row snapshot for an active,
        # exactly matched company. Returned workbook targets remain unverified.
        matched_company_id = uuid4()
        original_row.company_match = {
            "status": "matched",
            "candidate_id": str(matched_company_id),
        }
        original_row.match_status = "matched"
        session.get(ExternalImportActivation, original_row.id).active = True
        session.flush()
        statements = []

        def capture_statement(_conn, _cursor, statement, parameters, _context, _many):
            if "SELECT external_import_rows.id" in statement:
                statements.append((statement, parameters))

        sqlalchemy_event.listen(session.get_bind(), "before_cursor_execute", capture_statement)
        hint_result = company_hyperlink_hints(session, matched_company_id)
        sqlalchemy_event.remove(session.get_bind(), "before_cursor_execute", capture_statement)
        assert statements
        assert "external_import_rows.raw," not in statements[0][0]
        assert "external_import_batches.report," not in statements[0][0]
        assert 64 in statements[0][1].values()
        assert "octet_length" in statements[0][0]
        assert hint_result["urls"] == ["https://example.gov/permits/123"]
        assert hint_result["verification_status"] == "UNVERIFIED_SOURCE_HINT"
        assert len(hint_result["fingerprint"]) == 64
        assert (
            company_hyperlink_hints(session, matched_company_id)["fingerprint"]
            == (hint_result["fingerprint"])
        )
        session.get(ExternalImportActivation, original_row.id).active = False
        session.flush()

        exact_replay = stage_import(session, current_preview, actor="metadata-reviewer")
        assert exact_replay["metadata_reparse_receipts"] == reparsed["metadata_reparse_receipts"]
        assert session.scalar(select(func.count()).select_from(ExternalImportEvent)) == 2

        with pytest.raises(IntegrityError):
            with session.begin_nested():
                session.add(
                    ExternalImportEvent(
                        batch_id=original_row.batch_id,
                        action="metadata_reparsed",
                        actor="malformed-reviewer",
                        details={
                            "kind": "metadata_reparse",
                            "parser_revision": "invalid-hash-type",
                            "metadata_fingerprint": int("2" * 64),
                            "source_content_hash": "a" * 64,
                            "target_count": 1,
                            "hyperlink_hints": [],
                        },
                    )
                )
                session.flush()

        changed_display_row = current_preview.rows[0].model_copy(
            update={"raw": {**current_preview.rows[0].raw, "Project Name": "Changed display"}}
        )
        changed_display = current_preview.model_copy(
            update={"rows": [changed_display_row, *current_preview.rows[1:]]}
        )
        with pytest.raises(ValueError, match="changed original raw or normalized row data"):
            stage_import(session, changed_display, actor="metadata-reviewer")

        changed_target_raw = dict(current_preview.rows[0].raw)
        sidecar_key = next(
            key
            for key, value in changed_target_raw.items()
            if _is_hyperlink_sidecar(
                key,
                value,
                sidecar_keys=current_sidecar_keys.get(current_preview.rows[0].sheet, set()),
            )
        )
        changed_target_sidecar = dict(changed_target_raw[sidecar_key])
        changed_target_sidecar["Citation"] = "https://example.gov/permits/changed"
        changed_target_raw[sidecar_key] = changed_target_sidecar
        changed_target_row = current_preview.rows[0].model_copy(update={"raw": changed_target_raw})
        changed_target = current_preview.model_copy(
            update={"rows": [changed_target_row, *current_preview.rows[1:]]}
        )
        with pytest.raises(ValueError, match="parser revision produced different"):
            stage_import(session, changed_target, actor="metadata-reviewer")

        removed_target_row = current_preview.rows[0].model_copy(
            update={
                "raw": {
                    key: value
                    for key, value in current_preview.rows[0].raw.items()
                    if not _is_hyperlink_sidecar(
                        key,
                        value,
                        sidecar_keys=current_sidecar_keys.get(current_preview.rows[0].sheet, set()),
                    )
                }
            }
        )
        removed_targets = current_preview.model_copy(
            update={"rows": [removed_target_row, *current_preview.rows[1:]]}
        )
        with pytest.raises(ValueError, match="parser revision produced different"):
            stage_import(session, removed_targets, actor="metadata-reviewer")

        import gridlens.intelligence.external_imports as external_imports_module

        monkeypatch.setattr(
            external_imports_module, "HYPERLINK_METADATA_PARSER_REVISION", "future-v3"
        )
        future_preview = current_preview.model_copy(deep=True)
        future_report = dict(future_preview.report)
        future_report["hyperlink_metadata_parser_revision"] = "future-v3"
        future_preview.report = future_report
        future_replay = stage_import(session, future_preview, actor="metadata-reviewer")
        assert future_replay["metadata_reparse_receipts"] == reparsed["metadata_reparse_receipts"]
        assert session.scalar(select(func.count()).select_from(ExternalImportEvent)) == 2

        future_b_raw = dict(future_preview.rows[0].raw)
        future_b_key = next(
            key
            for key, value in future_b_raw.items()
            if _is_hyperlink_sidecar(
                key,
                value,
                sidecar_keys=current_sidecar_keys.get(future_preview.rows[0].sheet, set()),
            )
        )
        future_b_sidecar = dict(future_b_raw[future_b_key])
        future_b_sidecar["Citation"] = "https://example.gov/permits/future-b"
        future_b_raw[future_b_key] = future_b_sidecar
        future_b_row = future_preview.rows[0].model_copy(update={"raw": future_b_raw})
        future_b_preview = future_preview.model_copy(
            update={"rows": [future_b_row, *future_preview.rows[1:]]}
        )
        future_b_result = stage_import(session, future_b_preview, actor="metadata-reviewer")
        assert len(future_b_result["metadata_reparse_receipts"]) == 2
        assert session.scalar(select(func.count()).select_from(ExternalImportEvent)) == 3
        with pytest.raises(ValueError, match="parser revision produced different"):
            stage_import(session, future_preview, actor="metadata-reviewer")

        stored = session.scalar(select(ExternalImportRow))
        batch = session.get(ExternalImportBatch, stored.batch_id)
        assert stored.raw == original_raw
        assert stored.normalized == original_normalized
        assert batch.report == original_report
        assert session.scalar(select(func.count()).select_from(Project)) == 0
        assert session.scalar(select(func.count()).select_from(PowerNeed)) == 0
        assert session.scalar(select(func.count()).select_from(Evidence)) == 0
        assert not session.scalar(select(ExternalImportActivation.active))

        rolled_back = rollback_import(session, original["batch_id"], actor="reviewer")
        assert rolled_back["active_count"] == 0
        after_rollback_replay = stage_import(session, future_b_preview, actor="metadata-reviewer")
        assert after_rollback_replay["status"] == "rolled_back"
        assert after_rollback_replay["active_count"] == 0
        assert session.scalar(select(func.count()).select_from(ExternalImportEvent)) == 4
        history = import_history(session)
        assert history[0]["recovered_unverified_hyperlink_hint_count"] == 6

        oversized_hint = {
            "sheet": original_row.sheet,
            "row_number": original_row.row_number,
            "column": "oversized",
            "url": None,
            "target": "https://example.gov/" + "é" * 2048,
            "verification_status": "UNVERIFIED_SOURCE_HINT",
        }
        many_hints = [
            {
                "sheet": original_row.sheet,
                "row_number": original_row.row_number,
                "column": f"Citation {index}",
                "url": f"https://example.gov/permits/{index}",
                "target": f"https://example.gov/permits/{index}",
                "verification_status": "UNVERIFIED_SOURCE_HINT",
            }
            for index in range(150)
        ]
        many_hints.insert(0, oversized_hint)
        session.add(
            ExternalImportEvent(
                batch_id=original_row.batch_id,
                action="metadata_reparsed",
                actor="reader-bound-test",
                created_at=datetime.now(UTC) + timedelta(days=1),
                details={
                    "kind": "metadata_reparse",
                    "parser_revision": "reader-bound-test-v1",
                    "metadata_fingerprint": "f" * 64,
                    "source_content_hash": "a" * 64,
                    "target_count": len(many_hints),
                    "hyperlink_hints": many_hints,
                    "irrelevant_blob": "z" * 1_000_000,
                },
            )
        )
        session.get(ExternalImportActivation, original_row.id).active = True
        session.flush()
        bounded_statements = []

        def capture_bounded_query(_conn, _cursor, statement, parameters, _context, _many):
            if "SELECT external_import_rows.id" in statement:
                bounded_statements.append((statement, parameters))

        sqlalchemy_event.listen(session.get_bind(), "before_cursor_execute", capture_bounded_query)
        bounded = company_hyperlink_hints(session, matched_company_id)
        sqlalchemy_event.remove(session.get_bind(), "before_cursor_execute", capture_bounded_query)
        assert bounded["metadata_hints_may_be_truncated"] is True
        assert bounded["hints_returned"] == 64
        assert len(bounded["urls"]) == 8
        assert "https://example.gov/permits/0" in bounded["urls"]
        assert bounded["urls_truncated"] is True
        assert "octet_length" in bounded_statements[0][0]
        assert MAX_COMPANY_HYPERLINK_TARGET_BYTES in bounded_statements[0][1].values()


def test_fresh_hyperlink_stage_replay_does_not_add_a_reparse_event(import_db):
    with TemporaryDirectory() as temporary_directory:
        workbook_path = Path(temporary_directory) / "fresh-linked-projects.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Projects"
        sheet.append(["Project Name", "State", "Citation"])
        sheet.append(["Unmatched Project", "TX", "Link"])
        sheet["C2"].hyperlink = "https://example.gov/permits/fresh"
        workbook.save(workbook_path)
        preview = read_workbook(workbook_path)

    with import_db.begin() as session:
        first = stage_import(session, preview, actor="initial-importer")
        assert first["rows"][0]["unverified_hyperlink_hints"][0]["url"] == (
            "https://example.gov/permits/fresh"
        )
        replay = stage_import(session, preview, actor="initial-importer")
        assert replay["batch_id"] == first["batch_id"]
        assert replay["metadata_reparse_receipts"] == []
        assert session.scalar(select(func.count()).select_from(ExternalImportEvent)) == 1
        staged_row = session.scalar(select(ExternalImportRow))
        matched_company_id = uuid4()
        staged_row.match_status = "matched"
        staged_row.company_match = {
            "status": "matched",
            "candidate_id": str(matched_company_id),
        }
        session.get(ExternalImportActivation, staged_row.id).active = True
        session.flush()
        hints = company_hyperlink_hints(session, matched_company_id)
        assert hints["urls"] == ["https://example.gov/permits/fresh"]
        assert hints["verification_status"] == "UNVERIFIED_SOURCE_HINT"


def test_name_only_and_conflicting_project_signals_remain_proposals(import_db):
    with import_db.begin() as session:
        _seed_matchable(session)
        conflicting = Project(
            identity_key="test-site-2",
            name="North Campus",
            location="Austin, TX",
            source_id="dallas-monthly-permits",
            external_id="source-project-2",
        )
        session.add(conflicting)
        session.flush()
        preview = _preview(
            [
                ImportRow(
                    sheet="Sites",
                    row_number=2,
                    raw={},
                    normalized={
                        "company_name": "Acme Energy",
                        "project_name": "North Campus",
                        "city": "Austin",
                        "state": "TX",
                    },
                    errors=[],
                ),
                ImportRow(
                    sheet="Sites",
                    row_number=3,
                    raw={},
                    normalized={"project_name": "North Campus"},
                    errors=[],
                ),
            ],
            "d" * 64,
        )
        result = stage_import(session, preview, actor="reviewer")
        assert result["rows"][0]["company_match"]["status"] == "matched"
        assert result["rows"][0]["project_match"]["status"] == "ambiguous"
        assert result["rows"][1]["project_match"]["status"] == "unmatched"
        assert result["matched_count"] == 0
        assert session.scalar(select(func.count()).select_from(Project)) == 2


def test_read_only_reconciliation_reports_provenance_unknowns_and_no_promotion(import_db):
    with import_db.begin() as session:
        project, identity = _seed_matchable(session)
        preview = _preview(
            [
                ImportRow(
                    sheet="Workbook Sites",
                    row_number=17,
                    raw={"Developer": "Acme Energy", "Site": "North Campus", "Demand MW": 4},
                    normalized={
                        "company_name": "Acme Energy",
                        "project_name": "North Campus",
                        "city": "Austin",
                        "state": "TX",
                        "facility_demand_mw": 300,
                        "demand_mw": 4,
                    },
                    errors=[],
                )
            ],
            "e" * 64,
        )
        preview.report["worksheets"] = [
            {"name": "Workbook Sites", "mapped_columns": {"Demand MW": "demand_mw"}}
        ]
        before = session.scalar(select(func.count()).select_from(ExternalImportBatch))
        result = reconcile_preview(session, preview)
        row = result["rows"][0]
        assert result["read_only"] is True
        assert result["summary"]["exact_match"] == 1
        assert row["source"]["content_sha256"] == "e" * 64
        assert row["source"]["sheet"] == "Workbook Sites"
        assert row["source"]["row_number"] == 17
        assert row["company_match"]["candidate_id"] == str(identity.id)
        assert row["project_match"]["candidate_id"] == str(project.id)
        assert row["verification"] == "unverified_source_assertions"
        assert row["promotion"]["eligible"] is False
        assert any(item["field"] == "temporary_demand" for item in row["missing_data"])
        assert row["power_claims"]["facility_capacity_mw"] == 300
        assert row["power_claims"]["reported_demand_mw"] == 4
        assert row["power_claims"]["temporary_demand_mw"] is None
        assert row["power_claims"]["temporary_demand_basis"].startswith("UNKNOWN")
        assert row["power_claims"]["committed_supply_mw"] is None
        assert row["power_claims"]["residual_gap_mw"] is None
        assert any(item["field"] == "committed_supply" for item in row["missing_data"])
        assert session.scalar(select(func.count()).select_from(ExternalImportBatch)) == before
        assert session.scalar(select(func.count()).select_from(ExternalImportRow)) == 0
        assert session.scalar(select(func.count()).select_from(PowerNeed)) == 0


def test_read_only_reconciliation_bounds_large_duplicate_groups_and_truncated_input(import_db):
    with import_db.begin() as session:
        row_count = 25_000
        source_rows = [
            ImportRow(
                sheet="Projects",
                row_number=index + 2,
                raw={"Company": "Same Company"},
                normalized={},
                errors=[],
                row_kind="metadata",
            )
            for index in range(row_count)
        ]
        refs = [{"sheet": "Projects", "row_number": index + 2} for index in range(row_count)]
        preview = _preview(source_rows, "4" * 64)
        preview.report.update(
            {
                "truncated": True,
                "exact_duplicates": [{"key": "full_row_sha256", "rows": refs}],
            }
        )
        report = reconcile_preview(session, preview)
        assert report["complete"] is False
        assert report["processed_all_preview_rows"] is True
        assert report["row_details_truncated"] is True
        assert report["summary"]["rows"] == row_count
        assert report["summary"]["rows_included"] == 500
        assert report["summary"]["rows_omitted"] == row_count - 500
        assert report["summary"]["source_preview_truncated"] is True
        assert report["summary"]["duplicate_groups"] == 1
        assert len(report["duplicate_groups"][0]["sample_rows"]) == 10
        assert report["duplicate_groups"][0]["row_count"] == row_count
        assert len(report["rows"][0]["duplicate_hints"]) == 1
        assert report["rows"][0]["duplicate_hints"][0]["row_count"] == row_count
        assert "raw" not in report["rows"][0]["source"]


def test_explicit_activity_demand_header_is_separate_and_unverified(import_db):
    with import_db.begin() as session:
        result = reconcile_preview(
            session,
            ImportPreview(
                filename="temporary-demand.csv",
                content_hash="5" * 64,
                rows=[
                    ImportRow(
                        sheet="CSV",
                        row_number=2,
                        raw={"Temporary Demand MW": 7},
                        normalized={"demand_mw": 7},
                        errors=[],
                    )
                ],
                report={
                    "worksheets": [
                        {"name": "CSV", "mapped_columns": {"Temporary Demand MW": "demand_mw"}}
                    ]
                },
            ),
        )
        claims = result["rows"][0]["power_claims"]
        assert claims["temporary_demand_mw"] == 7
        assert claims["temporary_demand_basis"] == "unverified_activity_specific_source_assertion"
        assert claims["residual_gap_mw"] is None


def test_duplicate_project_names_resolve_by_distinct_physical_site(import_db):
    with import_db.begin() as session:
        first_project, identity = _seed_matchable(session)
        source = session.get(Source, "dallas-monthly-permits")
        second_project = Project(
            identity_key="test-site-houston",
            name=first_project.name,
            location="Houston, TX",
            source_id=source.id,
            external_id="source-project-houston",
        )
        observation = Observation(
            source_id=source.id,
            external_id="evidence-houston",
            title="North Campus construction permit",
            abstract="North Campus project site location: Houston, TX.",
            url="https://example.test/houston",
            content_hash="f" * 64,
        )
        session.add_all([second_project, observation])
        session.flush()
        evidence = Evidence(
            observation_id=observation.id,
            content_hash=observation.content_hash,
            title=observation.title,
            abstract=observation.abstract,
            url=observation.url,
            extraction_location="test excerpt",
            provenance={},
        )
        session.add(evidence)
        session.flush()
        retained = BuyingChainEvidence(
            project_id=second_project.id,
            retained_evidence_id=None,
            url=observation.url,
            title=observation.title,
            excerpt=observation.abstract,
            content_hash=observation.content_hash,
            extraction_location="test excerpt",
            published_at=None,
            observed_at=datetime.now(UTC),
            last_seen_at=datetime.now(UTC),
            project_match_reason="distinct physical site fixture",
        )
        session.add(retained)
        session.flush()
        session.add_all(
            [
                BuyingChainRun(
                    project_id=second_project.id,
                    input_hash="a" * 64,
                    evidence_ids=[str(retained.id)],
                    claim_ids=[],
                    policy_version="test",
                    status="complete",
                    documents_checked=1,
                    candidates_added=0,
                    unresolved=[],
                ),
                CompanyProjectMembership(
                    company_identity_id=identity.id,
                    project_id=second_project.id,
                    role="owner",
                    organization_name="Acme Energy",
                    source_id=source.id,
                    source_evidence_id=evidence.id,
                    source_version_hash=observation.content_hash,
                    decision_basis="test fixture",
                    operator_number="",
                ),
            ]
        )
        session.flush()
        result = stage_import(
            session,
            _preview(
                [
                    ImportRow(
                        sheet="Sites",
                        row_number=2,
                        raw={"Company": "Acme Energy", "Site": "North Campus"},
                        normalized={
                            "company_name": "Acme Energy",
                            "domain": "acme.example",
                            "project_name": "North Campus",
                            "city": "Houston",
                            "state": "TX",
                        },
                        errors=[],
                    )
                ],
                "1" * 64,
            ),
            actor="reviewer",
        )
        assert result["rows"][0]["project_match"]["status"] == "matched"
        assert result["rows"][0]["project_match"]["candidate_id"] == str(second_project.id)
        assert result["rows"][0]["project_match"]["candidate_id"] != str(first_project.id)


def test_reconciliation_surfaces_probable_project_without_auto_matching(import_db):
    with import_db.begin() as session:
        project, _identity = _seed_matchable(session)
        preview = _preview(
            [
                ImportRow(
                    sheet="Sites",
                    row_number=2,
                    raw={"Site": "Noth Campus", "City": "Austin"},
                    normalized={
                        "project_name": "Noth Campus",
                        "city": "Austin",
                        "state": "TX",
                    },
                    errors=[],
                )
            ],
            "2" * 64,
        )
        result = reconcile_preview(session, preview)
        row = result["rows"][0]
        assert result["summary"]["probable"] == 1
        assert row["status"] == "probable"
        assert row["project_match"]["status"] == "probable"
        assert row["project_match"]["candidate_id"] == str(project.id)
        assert "project name and city agree" in row["project_match"]["reasons"]
        assert row["promotion"]["eligible"] is False
        assert session.scalar(select(func.count()).select_from(ExternalImportRow)) == 0
        staged = stage_import(session, preview, actor="reviewer")
        assert staged["unmatched_count"] == 1
        assert staged["rows"][0]["project_match"]["status"] == "probable"
        assert staged["rows"][0]["project_match"]["candidate_id"] == str(project.id)
        assert staged["rows"][0]["active"] is False


def test_project_candidate_catalog_and_ambiguity_output_are_bounded(import_db):
    with import_db.begin() as session:
        source = Source(
            id="bounded-site-source",
            publisher="Synthetic",
            endpoint="https://bounded.example",
            family="construction-permits",
            geography="TX",
            industries=[],
            retrieval_method="bounded-test",
        )
        session.add(source)
        session.flush()
        projects = [
            Project(
                identity_key=f"bounded-site-{index}",
                name="North Campus",
                location="Austin, TX",
                region_state="TX",
                source_id=source.id,
                external_id=f"bounded-site-{index}",
            )
            for index in range(201)
        ]
        session.add_all(projects)
        session.flush()
        preview = _preview(
            [
                ImportRow(
                    sheet="Sites",
                    row_number=2,
                    raw={"Project": "North Campus"},
                    normalized={"project_name": "North Campus", "city": "Austin", "state": "TX"},
                    errors=[],
                ),
                ImportRow(
                    sheet="Sites",
                    row_number=3,
                    raw={"Project": "Noth Campus"},
                    normalized={"project_name": "Noth Campus", "city": "Austin", "state": "TX"},
                    errors=[],
                ),
            ],
            "6" * 64,
        )
        result = stage_import(session, preview, actor="reviewer")
        exact = result["rows"][0]["project_match"]
        probable = result["rows"][1]["project_match"]
        assert result["ambiguous_count"] == 2
        assert exact["status"] == "ambiguous"
        assert exact["candidate_count"] == 201
        assert len(exact["candidate_ids"]) <= 20
        assert probable["status"] == "ambiguous"
        assert probable["search_limited"] is True
        assert probable["candidate_count_lower_bound"] == 200
        assert len(probable["candidates"]) == 20
        assert probable["candidates_truncated"] is True
        assert all(not row["active"] for row in result["rows"])


def test_read_only_reconciliation_keeps_contact_assertions_unverified(import_db):
    with import_db.begin() as session:
        _project, identity = _seed_matchable(session)
        result = reconcile_preview(
            session,
            _preview(
                [
                    ImportRow(
                        sheet="Contacts",
                        row_number=9,
                        row_kind="contact",
                        raw={"Company": "Acme Energy", "Contact": "Jane Doe"},
                        normalized={
                            "company_name": "Acme Energy",
                            "domain": "acme.example",
                            "contact_name": "Jane Doe",
                            "email": "jane@acme.example",
                        },
                        errors=[],
                    )
                ],
                "3" * 64,
            ),
        )
        row = result["rows"][0]
        assert row["status"] == "exact_match"
        assert row["company_match"]["candidate_id"] == str(identity.id)
        assert row["project_match"]["status"] == "not_applicable"
        assert row["verification"] == "unverified_source_assertions"
        assert row["promotion"]["eligible"] is False
        assert session.scalar(select(func.count()).select_from(ExternalImportRow)) == 0


def test_state_conflict_vetoes_company_membership_project_match(import_db):
    with import_db.begin() as session:
        _project, _identity = _seed_matchable(session)
        preview = _preview(
            [
                ImportRow(
                    sheet="Sites",
                    row_number=2,
                    raw={},
                    normalized={
                        "company_name": "Acme Energy",
                        "domain": "acme.example",
                        "project_name": "North Campus",
                        "city": "Austin",
                        "state": "OK",
                    },
                    errors=[],
                ),
            ],
            "c" * 64,
        )
        result = stage_import(session, preview, actor="reviewer")
        assert result["rows"][0]["company_match"]["status"] == "matched"
        assert result["rows"][0]["project_match"]["status"] == "unmatched"
        assert result["matched_count"] == 0


def test_site_conflicts_veto_membership_and_unregistered_domain_fallback(import_db):
    with import_db.begin() as session:
        project, _identity = _seed_matchable(session)
        project.location = "100 First Avenue, Austin, TX"
        session.flush()
        rows = [
            ImportRow(
                sheet="Sites",
                row_number=2,
                raw={},
                normalized={
                    "company_name": "Acme Energy",
                    "domain": "acme.example",
                    "project_name": "North Campus",
                    "address": "999 Other Street",
                    "city": "Austin",
                    "state": "TX",
                },
                errors=[],
            ),
            ImportRow(
                sheet="Sites",
                row_number=3,
                raw={},
                normalized={
                    "company_name": "Acme Energy",
                    "domain": "acme.example",
                    "project_name": "North Campus",
                    "city": "Dallas",
                },
                errors=[],
            ),
            ImportRow(
                sheet="Sites",
                row_number=4,
                raw={},
                normalized={
                    "company_name": "Acme Energy",
                    "domain": "unrelated.example",
                    "project_name": "North Campus",
                    "address": "100 First Avenue",
                    "city": "Austin",
                    "state": "TX",
                },
                errors=[],
            ),
            ImportRow(
                sheet="Sites",
                row_number=5,
                raw={},
                normalized={
                    "company_name": "Acme Energy",
                    "project_name": "North Campus",
                    "address": "10 First Avenue",
                    "city": "Austin",
                    "state": "TX",
                },
                errors=[],
            ),
            ImportRow(
                sheet="Sites",
                row_number=6,
                raw={},
                normalized={
                    "company_name": "Acme Energy",
                    "domain": "acme.example",
                    "project_name": "North Campus",
                    "address": "100 First Avenue, Dallas, TX",
                },
                errors=[],
            ),
        ]
        result = stage_import(session, _preview(rows, "8" * 64), actor="reviewer")
        assert [row["project_match"]["status"] for row in result["rows"]] == [
            "unmatched",
            "unmatched",
            "matched",
            "unmatched",
            "unmatched",
        ]
        assert all(row["active"] is False for row in result["rows"])
        assert all(row["status"] == "unmatched" for row in result["rows"])
        assert result["active_count"] == 0


def test_company_matched_contact_only_row_is_active_and_unverified(import_db):
    with import_db.begin() as session:
        _seed_matchable(session)
        result = stage_import(
            session,
            _preview(
                [
                    ImportRow(
                        sheet="Contacts",
                        row_number=2,
                        row_kind="contact",
                        raw={"Company": "Acme Energy", "Email": "person@acme.example"},
                        normalized={
                            "company_name": "Acme Energy",
                            "email": "person@acme.example",
                            "contact_name": "Imported Person",
                        },
                        errors=[],
                    ),
                ],
                "e" * 64,
            ),
            actor="reviewer",
        )
        assert result["rows"][0]["company_match"]["status"] == "matched"
        assert result["rows"][0]["project_match"]["status"] == "not_applicable"
        assert result["rows"][0]["status"] == "matched"
        assert result["rows"][0]["active"] is True
        source_row = session.scalar(select(ExternalImportRow))
        assert source_row.normalized["contact_name"] == "Imported Person"
        assert source_row.normalized["email"] == "person@acme.example"


def test_summary_rows_are_retained_but_never_activated(import_db):
    with import_db.begin() as session:
        preview = _preview(
            [
                ImportRow(
                    sheet="Contacts",
                    row_number=2,
                    row_kind="summary",
                    raw={"Total records": 35},
                    normalized={},
                    errors=[],
                ),
            ],
            "b" * 64,
        )
        result = stage_import(session, preview, actor="reviewer")
        assert result["rows"][0]["status"] == "ignored"
        assert result["rows"][0]["row_kind"] == "summary"
        assert result["rows"][0]["active"] is False
        assert result["metrics"]["ignored_rows"] == 1


def test_stale_identifier_and_membership_evidence_never_drive_company_matches(import_db):
    with import_db.begin() as session:
        _seed_matchable(session)
        identifier = session.scalar(select(IdentityIdentifier))
        identifier.source_version_hash = "f" * 64
        membership = session.scalar(select(CompanyProjectMembership))
        membership.source_version_hash = "f" * 64
        result = stage_import(
            session,
            _preview(
                [
                    ImportRow(
                        sheet="Accounts",
                        row_number=2,
                        raw={},
                        normalized={"company_name": "Acme Energy", "domain": "acme.example"},
                        errors=[],
                    ),
                    ImportRow(
                        sheet="Accounts",
                        row_number=3,
                        raw={},
                        normalized={
                            "company_name": "Acme Energy",
                            "project_name": "North Campus",
                            "city": "Austin",
                            "state": "TX",
                        },
                        errors=[],
                    ),
                ],
                "9" * 64,
            ),
            actor="reviewer",
        )
        assert result["rows"][0]["company_match"]["status"] == "unmatched"
        assert result["rows"][1]["company_match"]["status"] == "unmatched"
        assert result["matched_count"] == 0


def test_staging_guard_rejects_arbitrary_database(import_db):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session

    production_shaped = create_engine("postgresql+psycopg://user:pass@localhost/gridlens")
    try:
        from gridlens.intelligence.external_imports import _assert_staging_session

        with Session(production_shaped) as session, pytest.raises(RuntimeError, match="restricted"):
            _assert_staging_session(session)
    finally:
        production_shaped.dispose()


def test_truncated_preview_is_rejected_before_any_write(import_db):
    with import_db.begin() as session:
        preview = _preview([], "f" * 64)
        preview.report["truncated"] = True
        with pytest.raises(ValueError, match="truncated"):
            stage_import(session, preview, actor="reviewer")
        assert session.scalar(select(func.count()).select_from(ExternalImportBatch)) == 0


def test_25k_rows_are_staged_in_bounded_chunks_and_return_bounded_summary(import_db):
    rows = [
        ImportRow(
            sheet="Accounts",
            row_number=index + 2,
            raw={"Account": f"Company {index}"},
            normalized={"company_name": f"Company {index}", "facility_demand_mw": "100"},
            errors=[],
        )
        for index in range(25_000)
    ]
    with import_db.begin() as session:
        result = stage_import(session, _preview(rows, "e" * 64), actor="reviewer")
        assert result["row_count"] == 25_000
        assert result["rows_truncated"] is True
        assert len(result["rows"]) == 100
        assert result["metrics"]["rows_with_power_assertions"] == 25_000
        assert session.scalar(select(func.count()).select_from(ExternalImportRow)) == 25_000


def test_same_bytes_metadata_reparse_changes_contact_digest_and_queues_new_job(import_db):
    from gridlens.intelligence.external_imports import company_hyperlink_hints

    with TemporaryDirectory() as temporary_directory:
        workbook_path = Path(temporary_directory) / "linked-accounts.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Accounts"
        sheet.append(["Company", "Domain", "Project", "City", "State", "Citation"])
        sheet.append(["Acme Energy", "acme.example", "North Campus", "Austin", "TX", "Link"])
        sheet.append(["Acme Energy", "acme.example", "North Campus", "Austin", "TX", "Link"])
        sheet.append(
            ["Unmatched Energy", "unmatched.example", "Unknown Site", "Austin", "TX", "Link"]
        )
        sheet["F2"].hyperlink = "https://acme.com/projects/north-campus"
        sheet["F3"].hyperlink = "https://inactive.com/acme"
        sheet["F4"].hyperlink = "https://unmatched.com/projects"
        workbook.save(workbook_path)

        current_preview = read_workbook(workbook_path)
        assert current_preview.report["hyperlink_metadata_parser_revision"] == (
            HYPERLINK_METADATA_PARSER_REVISION
        )
        sidecar_keys = _metadata_sidecar_keys(current_preview.report)
        legacy_report = {
            key: value
            for key, value in current_preview.report.items()
            if key not in {"hyperlink_metadata_parser_revision", "hyperlink_metadata_sidecar_keys"}
        }
        legacy_rows = [
            row.model_copy(
                update={
                    "raw": {
                        key: value
                        for key, value in row.raw.items()
                        if not _is_hyperlink_sidecar(
                            key,
                            value,
                            sidecar_keys=sidecar_keys.get(row.sheet, set()),
                        )
                    }
                }
            )
            for row in current_preview.rows
        ]
        legacy_preview = ImportPreview(
            filename=current_preview.filename,
            content_hash=current_preview.content_hash,
            rows=legacy_rows,
            report=legacy_report,
        )

    with import_db.begin() as session:
        project, company = _seed_matchable(session)
        project.region_status = "ELIGIBLE"
        session.add(
            PowerNeed(
                project_id=project.id,
                activity_key="construction",
                tier="Prospect",
                rationale="Staging planner fixture",
                required_mw="2-4 MW",
                start_window="2026-Q4",
            )
        )
        staged_v1 = stage_import(session, legacy_preview, actor="v1-importer")
        assert staged_v1["matched_count"] == 2
        first_rows = list(
            session.scalars(
                select(ExternalImportRow)
                .where(ExternalImportRow.batch_id == staged_v1["batch_id"])
                .order_by(ExternalImportRow.row_number)
            )
        )
        assert first_rows[0].raw["Citation"] == first_rows[1].raw["Citation"] == "Link"
        inactive_row = next(row for row in first_rows if row.row_number == 3)
        inactive_activation = session.get(ExternalImportActivation, inactive_row.id)
        assert inactive_activation is not None and inactive_activation.active
        inactive_activation.active = False

        assert company_hyperlink_hints(session, company.id)["urls"] == []
        before_source_digest = _research_source_digest(session, project, company.id)
        assert plan_company_contact_jobs(session, limit=10) == 1
        first_job = session.scalar(
            select(Job)
            .where(Job.kind == CONTACT_JOB_KIND, Job.target_id == company.id)
            .order_by(Job.created_at.desc(), Job.id.desc())
        )
        first_checkpoint = json.loads(first_job.checkpoint)
        first_checkpoint.update(status="complete", completed_at=datetime.now(UTC).isoformat())
        first_job.checkpoint = json.dumps(first_checkpoint)
        first_job.status = "succeeded"

        reparsed = stage_import(session, current_preview, actor="metadata-reparser")
        assert reparsed["batch_id"] == staged_v1["batch_id"]
        assert reparsed["recovered_unverified_hyperlink_hint_count"] == 3
        hint_result = company_hyperlink_hints(session, company.id)
        expected_url = "https://acme.com/projects/north-campus"
        assert hint_result["urls"] == [expected_url]
        assert hint_result["fingerprint"]
        assert _research_source_digest(session, project, company.id) != before_source_digest

        assert plan_company_contact_jobs(session, limit=10) == 1
        # Both jobs share the transaction timestamp; UUID ordering is random.
        new_job = session.scalars(
            select(Job).where(
                Job.kind == CONTACT_JOB_KIND,
                Job.target_id == company.id,
                Job.id != first_job.id,
            )
        ).one()
        new_checkpoint = json.loads(new_job.checkpoint)
        assert new_job.id != first_job.id
        assert new_checkpoint["research_digest"] != first_checkpoint["research_digest"]
        assert expected_url in new_checkpoint["payload"]["candidate_urls"]
