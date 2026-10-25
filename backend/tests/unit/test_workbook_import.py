from __future__ import annotations

import csv
import json
from datetime import date

from openpyxl import Workbook

from gridlens.ingestion.workbook_import import read_workbook


def test_multisheet_import_maps_mixed_rows_and_reports_duplicates(tmp_path):
    path = tmp_path / "leads.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Projects"
    first.append(["Data Center Name", "Developer", "State", "Power Capacity", "Completion Date"])
    first.append(["North Park", "Acme Power", "tx", 42, date(2027, 3, 1)])
    second = workbook.create_sheet("Contacts")
    second.append(["Organization", "Full Name", "Email Address", "Job Title"])
    second.append(["Acme Power", "Alex Example", "ALEX@example.com", "Director"])
    second.append(["Acme Power", "Other Person", "alex@example.com", "VP"])
    workbook.save(path)

    preview = read_workbook(path)

    assert len(preview.rows) == 3
    assert preview.rows[0].sheet == "Projects"
    assert preview.rows[0].row_number == 2
    assert preview.rows[0].normalized == {
        "project_name": "North Park",
        "company_name": "Acme Power",
        "state": "TX",
        "facility_demand_mw": 42,
        "facility_demand_mw_source_unit": "MW",
        "facility_demand_mw_unit_source": "canonical_default",
    }
    assert "required_date" not in preview.rows[0].normalized
    assert preview.rows[1].normalized["email"] == "alex@example.com"
    assert preview.report["exact_duplicates"] == []
    assert (
        json.loads(preview.model_dump_json())["rows"][0]["raw"]["Completion Date"]
        == "2027-03-01T00:00:00"
    )


def test_csv_mapping_and_bad_numeric_are_preview_errors(tmp_path):
    path = tmp_path / "custom.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Org", "Load (custom)", "Person"])
        writer.writerow(["Example LLC", "many", "Alex"])
    preview = read_workbook(
        path,
        column_mapping={
            "Org": "company_name",
            "Load (custom)": "demand_mw",
            "Person": "contact_name",
        },
    )
    row = preview.rows[0]
    assert row.normalized == {"company_name": "Example LLC", "contact_name": "Alex"}
    assert len(row.errors) == 1
    assert "demand_mw must be a scalar" in row.errors[0]


def test_formula_text_is_preserved_and_never_evaluated(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Company", "Email"])
    sheet.append(["=1+1", '=HYPERLINK("https://invalid","click")'])
    workbook.save(path)
    preview = read_workbook(path)
    assert preview.rows[0].raw["Company"] == "=1+1"
    assert preview.rows[0].normalized["company_name"] == "=1+1"
    assert preview.report["formulas_evaluated"] is False


def test_large_contact_csv_stays_bounded_and_keeps_all_rows(tmp_path):
    path = tmp_path / "contacts.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Company", "Name", "Email"])
        for index in range(25_000):
            writer.writerow([f"Company {index}", f"Person {index}", f"person{index}@example.com"])
    preview = read_workbook(path)
    assert len(preview.rows) == 25_000
    assert preview.report["truncated"] is False
    assert preview.rows[-1].normalized["contact_name"] == "Person 24999"


def test_full_row_duplicate_fingerprint_does_not_merge_distinct_projects(tmp_path):
    path = tmp_path / "projects.csv"
    path.write_text(
        "Developer,Project Name,State\nAcme,North,Texas\nAcme,South,Texas\nAcme,North,Texas\n",
        encoding="utf-8",
    )
    preview = read_workbook(path)
    assert preview.report["exact_duplicates"] == [
        {
            "key": "full_row_sha256",
            "value": preview.report["exact_duplicates"][0]["value"],
            "rows": [{"sheet": "CSV", "row_number": 2}, {"sheet": "CSV", "row_number": 4}],
        }
    ]
    assert preview.rows[0].normalized["state"] == "TX"
    assert len(preview.report["exact_duplicates"]) == 1


def test_repeated_headers_are_preserved_and_ambiguous_columns_reported(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Report export", None, None])
    sheet.append(["Developer", "Company", "Notes", "Notes", "Facility Power (MW)"])
    sheet.append(["Acme", "Acme LLC", "one", "two", 20])
    workbook.save(path)
    preview = read_workbook(path)
    row = preview.rows[0]
    assert row.row_number == 3
    assert row.normalized == {
        "facility_demand_mw": 20,
        "facility_demand_mw_source_unit": "MW",
        "facility_demand_mw_unit_source": "header",
    }
    assert "Developer" in row.raw and "Company" in row.raw
    assert row.raw["Notes"] == "one" and row.raw["Notes#2"] == "two"
    assert "ambiguous company_name columns" in row.errors[0]
    assert preview.report["worksheets"][0]["header_row_number"] == 2
    assert preview.report["worksheets"][0]["preamble_rows"] == 1


def test_conservative_capacity_date_domain_email_and_negative_values(tmp_path):
    path = tmp_path / "assertions.csv"
    path.write_text(
        "Project Name,Developer,Facility Power (MW),IT Capacity (MW),Completion Date,"
        "State,Website,Email,Supply MW\n"
        "North,Acme,-5,4,2028-01-01,New Mexico,https://www.acme.example/path,a bad,-1\n",
        encoding="utf-8",
    )
    preview = read_workbook(path)
    row = preview.rows[0]
    assert "facility_demand_mw" not in row.normalized
    assert row.normalized["it_capacity_mw"] == 4
    assert "required_date" not in row.normalized
    assert row.normalized["state"] == "NM"
    assert row.normalized["domain"] == "acme.example"
    assert "facility_demand_mw must be non-negative" in " ".join(row.errors)
    assert "supply_mw must be non-negative" in " ".join(row.errors)
    assert "email is invalid" in " ".join(row.errors)


def test_explicit_mapping_can_disambiguate_aliases(tmp_path):
    path = tmp_path / "mapped.csv"
    path.write_text("Company,Developer\nacme.example,Acme Inc\n", encoding="utf-8")
    preview = read_workbook(path, column_mapping={"Company": "domain"})
    assert preview.rows[0].normalized == {"domain": "acme.example", "company_name": "Acme Inc"}
    assert preview.rows[0].errors == []


def test_explicit_source_mapping_overrides_competing_automatic_alias(tmp_path):
    path = tmp_path / "owner-developer.csv"
    path.write_text(
        "Developer,Owner,Project Name\nAcme Development,Acme Holdings,North Site\n",
        encoding="utf-8",
    )
    preview = read_workbook(path, column_mapping={"Developer": "company_name"})
    assert preview.rows[0].normalized == {
        "company_name": "Acme Development",
        "project_name": "North Site",
    }
    assert preview.rows[0].errors == []


def test_summary_and_readme_rows_are_preserved_as_non_records(tmp_path):
    path = tmp_path / "tracker.xlsx"
    workbook = Workbook()
    project_sheet = workbook.active
    project_sheet.title = "All Projects"
    project_sheet.append(["Data Center Name", "Developer", "Address"])
    project_sheet.append(["North Site", "Acme", "Austin"])
    summary = workbook.create_sheet("State Summary")
    summary.append(["State", "Project Count"])
    summary.append(["TX", 158])
    readme = workbook.create_sheet("Read Me")
    readme.append(["Workbook Guide", "Value"])
    readme.append(["Source", "Research"])
    workbook.save(path)
    preview = read_workbook(path)
    assert preview.rows[0].row_kind == "project"
    assert preview.rows[1].row_kind == "summary"
    assert preview.rows[1].normalized == {}
    assert preview.rows[1].raw == {"State": "TX", "Project Count": 158}
    assert preview.rows[2].row_kind == "metadata"
    assert preview.rows[2].normalized == {}
    assert preview.report["recognized_dataset"][1]["dataset"] == "summary"


def test_power_units_convert_to_mw_with_source_unit_metadata(tmp_path):
    path = tmp_path / "units.csv"
    path.write_text(
        "Project Name,Facility Power (MW),IT Capacity (kW),Power Demand\n"
        "North,1.25 GW,500 kW,0.25\n",
        encoding="utf-8",
    )
    preview = read_workbook(path)
    row = preview.rows[0]
    assert row.normalized["facility_demand_mw"] == 1250
    assert row.normalized["facility_demand_mw_source_unit"] == "GW"
    assert row.normalized["facility_demand_mw_unit_source"] == "cell"
    assert row.normalized["it_capacity_mw"] == 0.5
    assert row.normalized["it_capacity_mw_source_unit"] == "KW"
    assert row.normalized["it_capacity_mw_unit_source"] == "cell"
    assert row.normalized["demand_mw"] == 0.25


def test_power_ranges_are_unverified_ranges_and_known_placeholders_are_missing(tmp_path):
    path = tmp_path / "ranges.csv"
    path.write_text(
        "Project Name,Power Capacity,IT Capacity (MW),Supply MW\n"
        "North,100 MW - 0.2 GW,Unknown,N/A\n"
        "South,TBD (MW),,Not disclosed\n",
        encoding="utf-8",
    )
    preview = read_workbook(path)
    row = preview.rows[0]
    assert "facility_demand_mw" not in row.normalized
    assert row.normalized["facility_demand_mw_range"] == {
        "min": 100,
        "max": 200,
        "unit": "MW",
        "basis": "imported_unverified",
        "source_units": ["MW", "GW"],
    }
    assert "it_capacity_mw" not in row.normalized
    assert "supply_mw" not in row.normalized
    assert row.errors == []
    assert preview.rows[1].normalized == {"project_name": "South"}
    assert preview.rows[1].errors == []


def test_arbitrary_power_prose_and_negative_values_remain_errors(tmp_path):
    path = tmp_path / "bad-power.csv"
    path.write_text(
        "Power Demand,Supply MW\n" + '"depends on phase",-0.5\n',
        encoding="utf-8",
    )
    row = read_workbook(path).rows[0]
    assert row.normalized == {}
    assert len(row.errors) == 2
    assert "other text is preserved in raw but not coerced" in row.errors[0]
    assert "supply_mw must be non-negative" in row.errors[1]
