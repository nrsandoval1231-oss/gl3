"""Bounded, non-mutating previews for externally supplied contact/project files."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from collections import defaultdict
from collections.abc import Iterator
from datetime import date, datetime, time
from difflib import SequenceMatcher
from itertools import chain
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict

CANONICAL_FIELDS = (
    "company_name",
    "domain",
    "project_name",
    "address",
    "city",
    "state",
    "external_id",
    "contact_name",
    "title",
    "email",
    "phone",
    "profile_url",
    "demand_mw",
    "facility_demand_mw",
    "it_capacity_mw",
    "supply_mw",
    "required_date",
    "duration_months",
)
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 300


class ImportRow(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sheet: str
    row_number: int
    raw: dict[str, Any]
    normalized: dict[str, Any]
    errors: list[str]
    row_kind: str = "record"


class ImportPreview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    filename: str
    content_hash: str
    rows: list[ImportRow]
    report: dict[str, Any]


_ALIASES: dict[str, tuple[str, ...]] = {
    "company_name": (
        "company",
        "company name",
        "organization",
        "organisation",
        "account",
        "account name",
        "employer",
        "business",
        "customer",
        "developer",
        "owner",
    ),
    "domain": ("domain", "website", "company website", "web site", "url", "company url"),
    "project_name": (
        "project",
        "project name",
        "site name",
        "facility",
        "data center name",
        "datacenter name",
    ),
    "address": ("address", "street", "street address", "site address", "location address"),
    "city": ("city", "site city", "project city"),
    "state": ("state", "state code", "site state", "project state"),
    "external_id": ("id", "external id", "record id", "project id", "contact id", "crm id"),
    "contact_name": ("contact", "contact name", "full name", "name", "person", "lead name"),
    "title": ("title", "job title", "position", "role"),
    "email": ("email", "email address", "e-mail", "project email", "business email"),
    "phone": ("phone", "telephone", "mobile", "phone number", "project phone"),
    "profile_url": ("profile url", "linkedin", "linkedin url", "linkedin profile", "profile link"),
    "demand_mw": (
        "temporary demand",
        "temporary demand mw",
        "temporary power demand",
        "temporary power demand mw",
        "rental power demand",
        "rental power demand mw",
        "construction power demand",
        "construction power demand mw",
        "commissioning power demand",
        "commissioning power demand mw",
        "demand",
        "demand mw",
        "demand kw",
        "demand gw",
        "load",
        "load mw",
        "load kw",
        "load gw",
        "power demand",
        "power demand mw",
        "power demand kw",
        "power demand gw",
        "required mw",
        "required kw",
        "required gw",
        "power requirement mw",
        "power requirement kw",
        "power requirement gw",
    ),
    "facility_demand_mw": (
        "facility power mw",
        "facility power kw",
        "facility power gw",
        "power capacity",
        "power capacity mw",
        "power capacity kw",
        "power capacity gw",
        "facility capacity mw",
    ),
    "it_capacity_mw": ("it capacity mw", "it capacity kw", "it capacity gw", "it capacity"),
    "supply_mw": (
        "supply mw",
        "supply kw",
        "supply gw",
        "available supply mw",
        "available supply kw",
        "available supply gw",
        "secured power mw",
        "secured power kw",
        "secured power gw",
        "utility supply mw",
        "utility supply kw",
        "utility supply gw",
    ),
    "required_date": (
        "required date",
        "power required date",
        "start date",
        "need date",
        "energization date",
        "target date",
    ),
    "duration_months": ("duration months", "duration", "term months", "months required"),
}

_MISSING_VALUES = {
    "n/a",
    "n.a.",
    "na",
    "none",
    "null",
    "unknown",
    "tbd",
    "not disclosed",
    "not available",
    "not provided",
    "undisclosed",
    "--",
    "—",
    "–",
    "-",
    "?",
}
_POWER_FIELDS = {"demand_mw", "facility_demand_mw", "it_capacity_mw", "supply_mw"}
_POWER_SCALAR = re.compile(
    r"^\s*(?P<number>[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(?P<unit>mw|gw|kw)?\s*$",
    re.IGNORECASE,
)
_POWER_RANGE = re.compile(
    r"^\s*(?P<low>[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(?P<unit1>mw|gw|kw)?\s*(?:-|–|—|to)\s*(?P<high>[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(?P<unit2>mw|gw|kw)?\s*$",
    re.IGNORECASE,
)
_UNIT_TO_MW = {"kw": 0.001, "mw": 1.0, "gw": 1000.0}

_STATE_NAMES = {
    "alabama": "AL",
    "alaska": "AK",
    "arizona": "AZ",
    "arkansas": "AR",
    "california": "CA",
    "colorado": "CO",
    "connecticut": "CT",
    "delaware": "DE",
    "florida": "FL",
    "georgia": "GA",
    "hawaii": "HI",
    "idaho": "ID",
    "illinois": "IL",
    "indiana": "IN",
    "iowa": "IA",
    "kansas": "KS",
    "kentucky": "KY",
    "louisiana": "LA",
    "maine": "ME",
    "maryland": "MD",
    "massachusetts": "MA",
    "michigan": "MI",
    "minnesota": "MN",
    "mississippi": "MS",
    "missouri": "MO",
    "montana": "MT",
    "nebraska": "NE",
    "nevada": "NV",
    "new hampshire": "NH",
    "new jersey": "NJ",
    "new mexico": "NM",
    "new york": "NY",
    "north carolina": "NC",
    "north dakota": "ND",
    "ohio": "OH",
    "oklahoma": "OK",
    "oregon": "OR",
    "pennsylvania": "PA",
    "rhode island": "RI",
    "south carolina": "SC",
    "south dakota": "SD",
    "tennessee": "TN",
    "texas": "TX",
    "utah": "UT",
    "vermont": "VT",
    "virginia": "VA",
    "washington": "WA",
    "west virginia": "WV",
    "wisconsin": "WI",
    "wyoming": "WY",
    "district of columbia": "DC",
}


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().casefold()).strip()


_ALIAS_TO_FIELD = {
    _header_key(alias): field for field, aliases in _ALIASES.items() for alias in (field, *aliases)
}


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, str | bool | int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, list | tuple):
        return [_json_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    return str(value)


def _normalized_value(field: str, value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value.casefold() in _MISSING_VALUES:
            return None
    if field == "email":
        email = str(value).strip().casefold()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            raise ValueError("email is invalid")
        return email
    if field == "domain":
        text = str(value).strip()
        parsed = urlparse(text if "://" in text else f"//{text}")
        host = (parsed.hostname or "").casefold().removeprefix("www.")
        if not host or "." not in host or any(ch.isspace() for ch in host):
            raise ValueError("domain must contain a valid host name")
        return host
    if field == "profile_url":
        return str(value).strip()
    if field == "state":
        state = str(value).strip().casefold()
        return _STATE_NAMES.get(state, state.upper())
    if field in {
        "demand_mw",
        "facility_demand_mw",
        "it_capacity_mw",
        "supply_mw",
        "duration_months",
    }:
        try:
            number = float(str(value).replace(",", "").strip())
            if not math.isfinite(number):
                raise ValueError
            if number < 0:
                raise ValueError
            return int(number) if number.is_integer() else number
        except (TypeError, ValueError):
            raise ValueError(
                f"{field} must be one non-negative number; units and ranges are not coerced"
            ) from None
    if field == "required_date" and isinstance(value, date | datetime):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _normalize_power_value(
    field: str, value: Any, header: Any
) -> tuple[Any, dict[str, Any] | None]:
    if value is None:
        return None, None
    if isinstance(value, str):
        missing_key = value.strip().casefold()
        missing_key = re.sub(r"\s*\(?\b(?:mw|gw|kw)\)?\s*$", "", missing_key).strip()
        if not missing_key or missing_key in _MISSING_VALUES:
            return None, None
    header_text = str(header or "")
    header_unit_match = re.search(r"\b(mw|gw|kw)\b", header_text, re.IGNORECASE)
    header_unit = header_unit_match.group(1).casefold() if header_unit_match else None
    source_text = str(value).strip()
    scalar = _POWER_SCALAR.fullmatch(source_text)
    range_match = _POWER_RANGE.fullmatch(source_text)

    def to_mw(number_text: str, unit: str) -> float:
        number = float(number_text.replace(",", ""))
        if not math.isfinite(number) or number < 0:
            raise ValueError(f"{field} must be non-negative")
        result = number * _UNIT_TO_MW[unit]
        if not math.isfinite(result):
            raise ValueError(f"{field} is outside the supported numeric range")
        return result

    if range_match:
        low_unit = (
            range_match.group("unit1") or range_match.group("unit2") or header_unit or "mw"
        ).casefold()
        high_unit = (
            range_match.group("unit2") or range_match.group("unit1") or header_unit or "mw"
        ).casefold()
        low = to_mw(range_match.group("low"), low_unit)
        high = to_mw(range_match.group("high"), high_unit)
        if low > high:
            raise ValueError(f"{field} range lower bound exceeds upper bound")
        return {
            "min": int(low) if low.is_integer() else low,
            "max": int(high) if high.is_integer() else high,
            "unit": "MW",
            "basis": "imported_unverified",
            "source_units": [low_unit.upper(), high_unit.upper()],
        }, {"source_unit": "range", "unit_source": "cell"}
    if scalar:
        unit_match = scalar.group("unit")
        unit = (unit_match or header_unit or "mw").casefold()
        number = to_mw(scalar.group("number"), unit)
        return (
            int(number) if number.is_integer() else number,
            {
                "source_unit": unit.upper(),
                "unit_source": "cell"
                if unit_match
                else "header"
                if header_unit
                else "canonical_default",
            },
        )
    raise ValueError(
        f"{field} must be a scalar with optional kW/MW/GW units or a numeric range; "
        "other text is preserved in raw but not coerced"
    )


def _mapping_for(
    headers: list[Any], column_mapping: dict[str, str] | None
) -> tuple[dict[int, str], list[str]]:
    explicit: dict[str, str] = {}
    for key, value in (column_mapping or {}).items():
        occurrence = re.fullmatch(r"(.*)#([2-9][0-9]*)", str(key).strip())
        lookup_key = (
            f"{_header_key(occurrence.group(1))} {occurrence.group(2)}"
            if occurrence
            else _header_key(key)
        )
        explicit[lookup_key] = value
    invalid = sorted({value for value in explicit.values() if value not in CANONICAL_FIELDS})
    if invalid:
        raise ValueError(f"Unknown canonical field(s) in column_mapping: {', '.join(invalid)}")
    result: dict[int, str] = {}
    explicit_indexes: set[int] = set()
    header_counts: dict[str, int] = defaultdict(int)
    for index, header in enumerate(headers):
        key = _header_key(header)
        header_counts[key] += 1
        explicit_field = explicit.get(f"{key} {header_counts[key]}") or explicit.get(key)
        field = explicit_field or _ALIAS_TO_FIELD.get(key)
        if field:
            result[index] = field
            if explicit_field:
                explicit_indexes.add(index)
    # Do not silently choose one of two source columns for the same canonical field.
    by_field: dict[str, list[int]] = defaultdict(list)
    for index, field in result.items():
        by_field[field].append(index)
    diagnostics = []
    for field, indexes in by_field.items():
        if len(indexes) > 1:
            explicit_for_field = [index for index in indexes if index in explicit_indexes]
            if len(explicit_for_field) == 1:
                selected = explicit_for_field[0]
                for index in indexes:
                    if index != selected:
                        result.pop(index)
                continue
            names = ", ".join(str(headers[index]) for index in indexes)
            diagnostics.append(f"ambiguous {field} columns: {names}; supply an explicit mapping")
            for index in indexes:
                result.pop(index)
    return result, diagnostics


def _header_score(values: tuple[Any, ...], column_mapping: dict[str, str] | None = None) -> int:
    explicit_headers = set()
    for header in column_mapping or {}:
        occurrence = re.fullmatch(r"(.*)#([2-9][0-9]*)", str(header).strip())
        explicit_headers.add(_header_key(occurrence.group(1) if occurrence else header))
    known_headers = set(_ALIAS_TO_FIELD) | explicit_headers
    return sum(_header_key(value) in known_headers for value in values if value is not None)


def _is_repeated_header_row(values: tuple[Any, ...], headers: list[Any]) -> bool:
    """Identify an echoed header row within a data range without interpreting it as a record."""
    if not headers or not any(header is not None and str(header).strip() for header in headers):
        return False
    for index, header in enumerate(headers):
        value = values[index] if index < len(values) else None
        if header is None or not str(header).strip():
            if value is not None and str(value).strip():
                return False
            continue
        if _header_key(value) != _header_key(header):
            return False
    return not any(value is not None and str(value).strip() for value in values[len(headers) :])


def _is_standalone_rows_footer(raw: dict[str, Any], project_name_headers: set[str]) -> bool:
    """Ignore the observed standalone footer label, but keep populated projects named Rows."""
    label_columns = {
        header
        for header in project_name_headers
        if str(raw.get(header) or "").strip().casefold() == "rows"
    }
    if not label_columns:
        return False
    return all(
        value is None or not str(value).strip()
        for header, value in raw.items()
        if header not in label_columns
    )


def _iter_sources(path: Path):
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            yield "CSV", enumerate(reader, start=1)
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook.worksheets:
                iterator = sheet.iter_rows(values_only=False)
                yield (
                    sheet.title,
                    (
                        (index, tuple(cell.value for cell in row))
                        for index, row in enumerate(iterator, start=1)
                    ),
                )
        finally:
            workbook.close()
    else:
        raise ValueError("Supported file types are .csv, .xlsx, and .xlsm")


def read_workbook(
    path: str | Path, *, column_mapping: dict[str, str] | None = None
) -> ImportPreview:
    """Read a bounded workbook or CSV into an immutable-value preview; never evaluates formulas."""
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds {MAX_FILE_BYTES} byte import limit")
    if source.suffix.casefold() in {".xlsx", ".xlsm"}:
        with zipfile.ZipFile(source) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > 1024 * 1024 * 1024:
                raise ValueError("Workbook exceeds the 1 GiB expanded-size import limit")
    digest = hashlib.sha256()
    with source.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)

    rows: list[ImportRow] = []
    sheets: list[dict[str, Any]] = []
    duplicate_index: dict[str, list[tuple[str, int]]] = defaultdict(list)
    probable_index: dict[str, list[tuple[str, int, str]]] = defaultdict(list)
    truncated = False
    for sheet_name, numbered_rows in _iter_sources(source):
        buffered = []
        for _ in range(20):
            try:
                buffered.append(next(numbered_rows))
            except StopIteration:
                break
        source_rows: Iterator[tuple[int, tuple[Any, ...]]]
        if buffered:
            header_position = max(
                range(len(buffered)),
                key=lambda index: (_header_score(buffered[index][1], column_mapping), -index),
            )
            header_row_number, header_values = buffered[header_position]
            headers = list(header_values)
            source_rows = iter(chain(buffered[header_position + 1 :], numbered_rows))
            preamble_count = header_position
        else:
            header_row_number, headers, source_rows, preamble_count = 1, [], iter(()), 0
        if len(headers) > MAX_COLUMNS:
            raise ValueError(f"Worksheet {sheet_name!r} exceeds {MAX_COLUMNS} columns")
        mapping, mapping_diagnostics = _mapping_for(headers, column_mapping)
        sheet_key = sheet_name.casefold().replace("_", " ").strip()
        if "read me" in sheet_key or "readme" in sheet_key or "guide" in sheet_key:
            row_kind = "metadata"
            recognized_dataset = "metadata"
        elif "summary" in sheet_key:
            row_kind = "summary"
            recognized_dataset = "summary"
        else:
            mapped_fields = set(mapping.values())
            has_project = "project_name" in mapped_fields
            has_contact = bool(mapped_fields & {"contact_name", "email", "phone", "profile_url"})
            if has_project and has_contact:
                row_kind = recognized_dataset = "mixed"
            elif has_project:
                row_kind = "project"
                recognized_dataset = "projects"
            elif has_contact:
                row_kind = "contact"
                recognized_dataset = "contacts"
            elif "company_name" in mapped_fields:
                row_kind = "company"
                recognized_dataset = "companies"
            else:
                row_kind = "record"
                recognized_dataset = "unclassified"
        base_headers = [
            str(header) if header is not None and str(header).strip() else f"column_{index + 1}"
            for index, header in enumerate(headers)
        ]
        reserved_headers = set(base_headers)
        actual_headers = {
            str(header) for header in headers if header is not None and str(header).strip()
        }
        used_headers: set[str] = set()
        duplicate_header_counts: dict[str, int] = defaultdict(int)
        raw_headers = []
        for index, base in enumerate(base_headers):
            is_actual_header = headers[index] is not None and bool(str(headers[index]).strip())
            if is_actual_header:
                duplicate_header_counts[base] += 1
                if duplicate_header_counts[base] == 1:
                    key = base
                else:
                    suffix = duplicate_header_counts[base]
                    key = f"{base}#{suffix}"
                    while key in reserved_headers or key in used_headers:
                        suffix += 1
                        key = f"{base}#{suffix}"
            else:
                key = base
                if key in actual_headers or key in used_headers:
                    suffix = 2
                    key = f"{base}#{suffix}"
                    while key in reserved_headers or key in used_headers:
                        suffix += 1
                        key = f"{base}#{suffix}"
            used_headers.add(key)
            raw_headers.append(key)
        start_count = len(rows)
        sheets.append(
            {
                "name": sheet_name,
                "header_row_number": header_row_number,
                "preamble_rows": preamble_count,
                "headers": [_json_value(h) for h in headers],
                "mapped_columns": {raw_headers[i]: field for i, field in mapping.items()},
                "mapping_diagnostics": mapping_diagnostics,
                "recognized_dataset": recognized_dataset,
                "row_kind": row_kind,
            }
        )
        for row_number, values in source_rows:
            if len(rows) >= MAX_ROWS:
                truncated = True
                break
            raw = {
                raw_headers[i]: _json_value(values[i]) if i < len(values) else None
                for i in range(len(headers))
            }
            for index in range(len(headers), len(values)):
                base = f"_extra_column_{index + 1}"
                key = base
                suffix = 2
                while key in reserved_headers or key in used_headers:
                    key = f"{base}#{suffix}"
                    suffix += 1
                raw[key] = _json_value(values[index])
            if not any(value is not None and value != "" for value in raw.values()):
                continue
            project_name_headers = {
                raw_headers[index] for index, field in mapping.items() if field == "project_name"
            }
            is_footer = _is_standalone_rows_footer(raw, project_name_headers)
            item_row_kind = (
                "metadata" if _is_repeated_header_row(values, headers) or is_footer else row_kind
            )
            normalized: dict[str, Any] = {}
            errors: list[str] = []
            for index, field in (
                mapping.items() if item_row_kind not in {"summary", "metadata"} else ()
            ):
                value = values[index] if index < len(values) else None
                try:
                    if field in _POWER_FIELDS:
                        cleaned, unit_metadata = _normalize_power_value(
                            field, value, headers[index]
                        )
                        if cleaned is not None:
                            if isinstance(cleaned, dict):
                                normalized[f"{field}_range"] = cleaned
                            else:
                                normalized[field] = cleaned
                            if unit_metadata:
                                normalized[f"{field}_source_unit"] = unit_metadata["source_unit"]
                                normalized[f"{field}_unit_source"] = unit_metadata["unit_source"]
                        continue
                    cleaned = _normalized_value(field, value)
                    if cleaned is not None:
                        normalized[field] = cleaned
                except ValueError as exc:
                    errors.append(f"{headers[index]}: {exc}")
            item = ImportRow(
                sheet=sheet_name,
                row_number=row_number,
                raw=raw,
                normalized=normalized,
                errors=errors,
                row_kind=item_row_kind,
            )
            rows.append(item)
            if mapping_diagnostics:
                item.errors.extend(mapping_diagnostics)
            fingerprint = hashlib.sha256(
                json.dumps(raw, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode(
                    "utf-8"
                )
            ).hexdigest()
            duplicate_index[fingerprint].append((sheet_name, row_number))
            company = str(normalized.get("company_name", "")).strip().casefold()
            if company:
                probable_index.setdefault(company, []).append((sheet_name, row_number, company))
        sheets[-1]["rows_processed"] = len(rows) - start_count
        if truncated:
            break

    exact = [
        {
            "key": "full_row_sha256",
            "value": fingerprint,
            "rows": [{"sheet": sheet, "row_number": row} for sheet, row in refs],
        }
        for fingerprint, refs in duplicate_index.items()
        if len(refs) > 1
    ]
    companies = list(probable_index.items())
    probable: list[dict[str, Any]] = []
    candidates: set[tuple[str, str]] = set()
    ordered_companies = sorted(companies)
    # Similar names are adjacent lexicographically in common typo cases. A
    # bounded neighborhood avoids quadratic work on large contact exports.
    for index, (left, _) in enumerate(ordered_companies):
        for right, _ in ordered_companies[index + 1 : index + 9]:
            if left[:3] == right[:3] and abs(len(left) - len(right)) <= max(
                2, int(max(len(left), len(right)) * 0.12)
            ):
                candidates.add((left, right))
    for left, right in sorted(candidates):
        ratio = SequenceMatcher(None, left, right).ratio()
        if ratio >= 0.88:
            refs = probable_index[left] + probable_index[right]
            probable.append(
                {
                    "left": left,
                    "right": right,
                    "similarity": round(ratio, 3),
                    "rows": [{"sheet": s, "row_number": n} for s, n, _ in refs],
                }
            )
    return ImportPreview(
        filename=source.name,
        content_hash=digest.hexdigest(),
        rows=rows,
        report={
            "worksheets": sheets,
            "recognized_dataset": [
                {
                    "sheet": sheet["name"],
                    "dataset": sheet["recognized_dataset"],
                    "row_kind": sheet["row_kind"],
                }
                for sheet in sheets
            ],
            "rows_processed": len(rows),
            "rows_with_errors": sum(bool(row.errors) for row in rows),
            "exact_duplicates": exact,
            "probable_duplicates": probable,
            "truncated": truncated,
            "limit_reached": truncated,
            "limits": {"file_bytes": MAX_FILE_BYTES, "rows": MAX_ROWS, "columns": MAX_COLUMNS},
            "row_limit": MAX_ROWS,
            "file_size_limit_bytes": MAX_FILE_BYTES,
            "formulas_evaluated": False,
        },
    )
